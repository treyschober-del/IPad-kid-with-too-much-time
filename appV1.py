import tkinter as tk
from tkinter import ttk, messagebox
import configparser
from openpyxl import load_workbook
import os
import time
import threading

DARK_BG = "#121212"
DARK_FG = "#f0f0f0"
ACCENT = "#3f51b5"
ERROR_COLOR = "#b00020"

STATS = ["Body", "Mind", "Soul", "Unconscious"]
STAT_CELL_MAP = {
    "Body": "W10",
    "Mind": "W11",
    "Soul": "W12",
    "Unconscious": "W13",
}
HEALTH_CELL_MAP = {
    "Body": "X10",
    "Mind": "X11",
    "Soul": "X12",
    "Unconscious": "X13",
}
NAME_CELL = "V5"
CHARGES_CELL = "B2"


class Tooltip:
    def __init__(self, widget, text):
        self.widget = widget
        self.text = text
        self.tipwindow = None
        widget.bind("<Enter>", self.show_tip)
        widget.bind("<Leave>", self.hide_tip)

    def show_tip(self, event=None):
        if self.tipwindow or not self.text:
            return
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + 20
        self.tipwindow = tw = tk.Toplevel(self.widget)
        tw.wm_overrideredirect(True)
        tw.configure(bg=DARK_BG)
        label = tk.Label(
            tw,
            text=self.text,
            justify=tk.LEFT,
            bg=DARK_BG,
            fg=DARK_FG,
            relief=tk.SOLID,
            borderwidth=1,
            padx=5,
            pady=3,
            wraplength=300,
        )
        label.pack()

        tw.wm_geometry(f"+{x}+{y}")

    def hide_tip(self, event=None):
        if self.tipwindow:
            self.tipwindow.destroy()
            self.tipwindow = None


class DnDApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Homebrew Campaign Manager")
        self.configure(bg=DARK_BG)
        self.style = ttk.Style(self)
        self._setup_style()
        self.attributes("-fullscreen", True)
        self.bind("<Escape>", lambda e: self.attributes("-fullscreen", False))


        self.config = self._load_config()
        self.stats_wb, self.stats_ws = self._load_workbook(self.config["files"]["stats_file"])
        self.abilities_wb, self.abilities_ws = self._load_workbook(self.config["files"]["abilities_file"])

        # --- Notes sheet setup ---
        if "Notes" in self.stats_wb.sheetnames:
            self.notes_ws = self.stats_wb["Notes"]
        else:
            self.notes_ws = self.stats_wb.create_sheet("Notes")
            # Optional: put a header
            self.notes_ws["A1"] = "Campaign Notes"
            self.stats_wb.save(self.config["files"]["stats_file"])

        self.stat_levels = {}
        self.stat_health = {}
        self.max_health = {}
        self.charges = 0

        self._load_stats()
        self._load_charges()

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)

        self.stats_frame = tk.Frame(self.notebook, bg=DARK_BG)
        self.abilities_frame = tk.Frame(self.notebook, bg=DARK_BG)
        self.new_item_frame = tk.Frame(self.notebook, bg=DARK_BG)
        self.notes_frame = tk.Frame(self.notebook, bg=DARK_BG)

        self.notebook.add(self.stats_frame, text="STATS & HEALTH")
        self.notebook.add(self.abilities_frame, text="ABILITIES")
        self.notebook.add(self.new_item_frame, text="NEW ITEM")
        self.notebook.add(self.notes_frame, text="NOTES")

        self._build_notes_tab()

        self.items = []          # list of dicts for all loaded items
        self.active_items = set()  # names of active items
        self.item_mods = {"body": 0, "mind": 0, "soul": 0, "unconscious": 0}

        self.charges_var = tk.IntVar(value=self.charges)
        self._build_global_header()

        self.items_frame = tk.Frame(self.notebook, bg=DARK_BG)
        self.notebook.add(self.items_frame, text="ITEMS")

        self._build_items_tab()

        self.health_bars = {}
        self.health_labels = {}
        self._build_stats_tab()
        # Red health bar style
        self.style.configure(
            "Red.Horizontal.TProgressbar",
            troughcolor=DARK_BG,
            background="#ff0000",   # bright red
            bordercolor=DARK_BG,
            lightcolor="#ff0000",
            darkcolor="#ff0000"
        )


        self.abilities_buttons = []
        self._build_abilities_tab()

        self._build_new_item_tab()

        self.invalid_overlay = None

        self.after(100, self._handle_name_and_intro)

    # ---------- Setup & Data ----------
    def _build_items_tab(self):
        for w in self.items_frame.winfo_children():
            w.destroy()

        header = tk.Frame(self.items_frame, bg=DARK_BG)
        header.pack(fill="x", pady=10)

        tk.Label(
            header,
            text="Inventory (max 5 active items)",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 14, "bold"),
        ).pack(side="left", padx=10)

        tk.Button(
            header,
            text="Add Item (.txt)",
            command=self._add_item_from_file,
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(side="right", padx=10)

        # Scrollable list
        container = tk.Frame(self.items_frame, bg=DARK_BG)
        container.pack(fill="both", expand=True, padx=10, pady=10)

        canvas = tk.Canvas(container, bg=DARK_BG, highlightthickness=0)
        scrollbar = tk.Scrollbar(container, orient="vertical", command=canvas.yview)
        self.items_list_frame = tk.Frame(canvas, bg=DARK_BG)

        self.items_list_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=self.items_list_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Render existing items (if any)
        for item in self.items:
            self._render_item_row(item)


    def _add_item_from_file(self):
        from tkinter import filedialog

        path = filedialog.askopenfilename(filetypes=[("Item files", "*.txt")])
        if not path:
            return

        item = self._load_item_from_txt(path)
        if not item:
            return

        self.items.append(item)
        self._render_item_row(item)

    def _load_item_from_txt(self, txt_path):
        folder = os.path.dirname(txt_path)
        data = {}
        try:
            with open(txt_path, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    parts = line.split("\t", 1)
                    if len(parts) == 2:
                        k, v = parts
                    else:
                        if ":" in line:
                            k, v = line.split(":", 1)
                        else:
                            continue
                    data[k.strip().lower()] = v.strip()
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read item file:\n{e}")
            return None

        name = data.get("name")
        if not name:
            messagebox.showerror("Error", "Item file missing 'name'.")
            return None

        # Basic stat mods (default 0)
        def get_int(key):
            try:
                return int(data.get(key, "0"))
            except ValueError:
                return 0

        item = {
            "name": name,
            "folder": folder,
            "body": get_int("body"),
            "mind": get_int("mind"),
            "soul": get_int("soul"),
            "unconscious": get_int("unconscious"),
            "extra": {k: v for k, v in data.items() if k not in ["name", "body", "mind", "soul", "unconscious"]},
            "image": None,
        }

        # Find a PNG in the folder
        for fname in os.listdir(folder):
            if fname.lower().endswith(".png"):
                item["image"] = os.path.join(folder, fname)
                break

        return item

    def _render_item_row(self, item):
        row = tk.Frame(self.items_list_frame, bg=DARK_BG)
        row.pack(fill="x", pady=5)

        # Optional image thumbnail
        img_label = tk.Label(row, bg=DARK_BG)
        img_label.pack(side="left", padx=5)

        if item.get("image"):
            try:
                from PIL import Image, ImageTk
                img = Image.open(item["image"])
                img.thumbnail((48, 48))
                photo = ImageTk.PhotoImage(img)
                img_label.image = photo  # keep reference
                img_label.configure(image=photo)
            except Exception:
                pass

        # Name + stats
        text = f"{item['name']}  (B:{item['body']} M:{item['mind']} S:{item['soul']} U:{item['unconscious']})"
        tk.Label(row, text=text, bg=DARK_BG, fg=DARK_FG, anchor="w").pack(side="left", padx=5)

        var = tk.BooleanVar(value=False)

        def on_toggle():
            if var.get():
                # Enforce max 5 active
                if len(self.active_items) >= 5:
                    messagebox.showwarning("Limit reached", "You can only have 5 active items.")
                    var.set(False)
                    return
                self.active_items.add(item["name"])
            else:
                self.active_items.discard(item["name"])
            self._recompute_item_mods()

        tk.Checkbutton(
            row,
            text="Active",
            variable=var,
            command=on_toggle,
            bg=DARK_BG,
            fg=DARK_FG,
            selectcolor=DARK_BG,
        ).pack(side="right", padx=5)


    def _setup_style(self):
        self.style.theme_use("clam")
        self.style.configure("TNotebook", background=DARK_BG, borderwidth=0)
        self.style.configure("TNotebook.Tab", background="#1e1e1e", foreground=DARK_FG)
        self.style.map("TNotebook.Tab", background=[("selected", ACCENT)])
        self.style.configure("TButton", background="#1e1e1e", foreground=DARK_FG)
        self.style.map("TButton", background=[("active", ACCENT)])

    def _load_config(self):
        config = configparser.ConfigParser()
        if not os.path.exists("settings.ini"):
            messagebox.showerror("Error", "settings.ini not found.")
            self.destroy()
            raise SystemExit
        config.read("settings.ini")
        return config

    def _load_workbook(self, path):
        if not os.path.exists(path):
            messagebox.showerror("Error", f"File not found:\n{path}")
            self.destroy()
            raise SystemExit
        wb = load_workbook(path)
        ws = wb.active
        return wb, ws

    def _build_notes_tab(self):
        tk.Label(
            self.notes_frame,
            text="Campaign Notes",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 14, "bold"),
        ).pack(pady=10)

        container = tk.Frame(self.notes_frame, bg=DARK_BG)
        container.pack(fill="both", expand=True, padx=10, pady=10)

        scrollbar = tk.Scrollbar(container)
        scrollbar.pack(side="right", fill="y")

        self.notes_text = tk.Text(
            container,
            bg="#1e1e1e",
            fg=DARK_FG,
            insertbackground=DARK_FG,
            wrap="word",
            yscrollcommand=scrollbar.set,
        )
        self.notes_text.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.notes_text.yview)

        # Load existing notes from A2 (A1 is optional header)
        existing = self.notes_ws["A2"].value or ""
        self.notes_text.insert("1.0", existing)

        tk.Button(
            self.notes_frame,
            text="Save Notes",
            command=self._save_notes,
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(pady=10)

    def _save_notes(self):
        content = self.notes_text.get("1.0", "end-1c")
        self.notes_ws["A2"] = content
        self.stats_wb.save(self.config["files"]["stats_file"])


    def _cell_value(self, ws, cell):
        return ws[cell].value

    def _set_cell_value(self, ws, cell, value):
        ws[cell].value = value

    def _save_all(self):
        self.stats_wb.save(self.config["files"]["stats_file"])
        self.abilities_wb.save(self.config["files"]["abilities_file"])

    def _load_stats(self):
        self.base_health = {}
        for stat in STATS:
            level_cell = STAT_CELL_MAP[stat]
            level = self._cell_value(self.stats_ws, level_cell) or 0
            try:
                level = int(level)
            except ValueError:
                level = 0
            level = max(0, min(4, level))
            self.stat_levels[stat] = level

            max_hp = 0
            if level > 0:
                max_hp = min(4, level) + 1  # 1->2, 2->3, 3->4, 4->5
            self.max_health[stat] = max_hp

            if stat == "Unconscious":
                self.stat_health[stat] = 0
                self.base_health[stat] = 0
                continue

            health_cell = HEALTH_CELL_MAP[stat]
            current = self._cell_value(self.stats_ws, health_cell)
            if current is None:
                current = max_hp
            try:
                current = int(current)
            except ValueError:
                current = max_hp
            current = max(0, min(max_hp, current))
            self.base_health[stat] = current
            self.stat_health[stat] = current 

    def _load_charges(self):
        val = self._cell_value(self.abilities_ws, CHARGES_CELL) or 0
        try:
            self.charges = int(val)
        except ValueError:
            self.charges = 0

    # ---------- Intro / Name ----------
    def _recompute_item_mods(self):
        # Reset mods
        self.item_mods = {"body": 0, "mind": 0, "soul": 0, "unconscious": 0}
        for item in self.items:
            if item["name"] in self.active_items:
                self.item_mods["body"] += item["body"]
                self.item_mods["mind"] += item["mind"]
                self.item_mods["soul"] += item["soul"]
                self.item_mods["unconscious"] += item["unconscious"]
        self._recompute_display_health()
        self._update_abilities_state()

    def _recompute_display_health(self):
        for stat in ["Body", "Mind", "Soul"]:
            base = self.base_health.get(stat, 0)
            key = stat.lower()
            bonus = self.item_mods.get(key, 0)
            max_hp = self.max_health[stat]
            effective = max(0, min(max_hp, base + bonus))
            self.stat_health[stat] = effective
            if stat in self.health_bars:
                self.health_bars[stat].set(effective)
            if stat in self.health_labels:
                self.health_labels[stat].set(f"{effective}/{max_hp}")


    def _handle_name_and_intro(self):
        name = self._cell_value(self.stats_ws, NAME_CELL)
        if not name:
            self._prompt_for_name()
        else:
            self._welcome_back(name)

    def _prompt_for_name(self):
        dialog = tk.Toplevel(self)
        dialog.title("Enter Character Name")
        dialog.configure(bg=DARK_BG)
        dialog.grab_set()

        tk.Label(dialog, text="Enter your character name:", bg=DARK_BG, fg=DARK_FG).pack(padx=20, pady=10)
        entry = tk.Entry(dialog, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG)
        entry.pack(padx=20, pady=5)
        entry.focus_set()

        def save_name():
            name = entry.get().strip()
            if not name:
                messagebox.showwarning("Warning", "Name cannot be empty.")
                return
            self._set_cell_value(self.stats_ws, NAME_CELL, name)
            self._save_all()
            dialog.destroy()
            self._welcome_back(name)

        tk.Button(dialog, text="OK", command=save_name, bg="#1e1e1e", fg=DARK_FG).pack(pady=10)

    def _welcome_back(self, name):
        overlay = tk.Toplevel(self)
        overlay.configure(bg=DARK_BG)
        overlay.attributes("-fullscreen", True)

        # Centered label
        label = tk.Label(
            overlay,
            text=f"Welcome back, {name}!",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 32, "bold")
        )
        label.pack(expand=True)  # <--- This keeps it centered forever

        # Auto-close after 3 seconds
        overlay.after(3000, overlay.destroy)



        def close_after_delay():
            time.sleep(3)
            try:
                overlay.destroy()
            except tk.TclError:
                pass

        threading.Thread(target=close_after_delay, daemon=True).start()

    # ---------- Global Header (Charges) ----------

    def _build_global_header(self):
        header = tk.Frame(self, bg=DARK_BG)
        header.place(relx=1.0, rely=0.0, anchor="ne")

        tk.Label(
            header,
            text="Charges Available:",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 10, "bold"),
        ).pack(side="left", padx=(0, 5))

        tk.Label(
            header,
            textvariable=self.charges_var,
            bg=DARK_BG,
            fg=ACCENT,
            font=("Papyrus", 10, "bold"),
        ).pack(side="left")

    # ---------- Stats & Health Tab ----------

    def _build_stats_tab(self):
        tk.Label(
            self.stats_frame,
            text="Stats & Health",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 14, "bold"),
        ).pack(pady=10)

        tab_container = tk.Frame(self.stats_frame, bg=DARK_BG)
        tab_container.pack(expand=True)

        # Save & Quit button (bottom-right corner)
        def save_and_quit():
            self._save_all()
            self.destroy()

        tk.Button(
            text="Save & Quit",
            command=save_and_quit,
            bg="white",
            fg="black",
            font=("Papyrus", 16, "bold")
        ).place(relx=0.98, rely=0.98, anchor="se")


        for stat in STATS:
            row = tk.Frame(tab_container, bg=DARK_BG)
            row.pack(fill="x", pady=5)

            tk.Label(row, text=stat, bg=DARK_BG, fg=DARK_FG, width=12, anchor="w").pack(side="left")

            if stat == "Unconscious":
                tk.Label(
                    row,
                    text="(No HP)",
                    bg=DARK_BG,
                    fg="#888888",
                    anchor="w",
                ).pack(side="left", padx=5)
                continue

            max_hp = self.max_health[stat]
            current_hp = self.stat_health[stat]

            bar_var = tk.IntVar(value=current_hp)
            bar = ttk.Progressbar(
                row,
                maximum=max_hp if max_hp > 0 else 1,
                variable=bar_var,
                length=200,
                style="Red.Horizontal.TProgressbar"   # <-- apply red style
            )

            bar.pack(side="left", padx=5)
            self.health_bars[stat] = bar_var

            label_var = tk.StringVar(value=f"{current_hp}/{max_hp}")
            lbl = tk.Label(row, textvariable=label_var, bg=DARK_BG, fg=DARK_FG, width=8)
            lbl.pack(side="left", padx=5)
            self.health_labels[stat] = label_var

            btn_frame = tk.Frame(row, bg=DARK_BG)
            btn_frame.pack(side="left", padx=5)

            tk.Button(
                btn_frame,
                text="-",
                width=2,
                command=lambda s=stat: self._change_health(s, -1),
                bg="#1e1e1e",
                fg=DARK_FG,
            ).pack(side="left", padx=2)
            tk.Button(
                btn_frame,
                text="+",
                width=2,
                command=lambda s=stat: self._change_health(s, 1),
                bg="#1e1e1e",
                fg=DARK_FG,
            ).pack(side="left", padx=2)

    def _change_health(self, stat, delta):
        if stat == "Unconscious":
            return
        max_hp = self.max_health[stat]
        base = self.base_health.get(stat, 0)
        new_base = max(0, min(max_hp, base + delta))
        self.base_health[stat] = new_base

        # Save base to Excel
        health_cell = HEALTH_CELL_MAP[stat]
        self._set_cell_value(self.stats_ws, health_cell, new_base)
        self._save_all()

        # Recompute effective (base + items)
        self._recompute_display_health()
        self._update_abilities_state()
        self._check_invalid_state()


    def _check_invalid_state(self):
        zero_stats = [s for s in ["Body", "Mind", "Soul"] if self.stat_health.get(s, 0) == 0]
        if len(zero_stats) >= 2:
            self._show_invalid_screen(zero_stats)

    def _show_invalid_screen(self, zero_stats):
        if self.invalid_overlay and tk.Toplevel.winfo_exists(self.invalid_overlay):
            return

        self.invalid_overlay = tk.Toplevel(self)
        self.invalid_overlay.configure(bg=DARK_BG)
        self.invalid_overlay.attributes("-fullscreen", True)

        container = tk.Frame(self.invalid_overlay, bg=DARK_BG)
        container.pack(expand=True)  # <--- centers everything

        tk.Label(
            container,
            text="INVALID",
            bg=DARK_BG,
            fg=ERROR_COLOR,
            font=("Papyrus", 48, "bold")
        ).pack(pady=20)

        tk.Label(
            container,
            text="You have lost two or more stats.",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 20)
        ).pack(pady=10)

        btn_frame = tk.Frame(container, bg=DARK_BG)
        btn_frame.pack(pady=20)


        tk.Button(
            btn_frame,
            text="Friends have saved me",
            command=lambda: self._revive(zero_stats, "friends"),
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(side="left", padx=10)

        tk.Button(
            btn_frame,
            text="Left to rot",
            command=lambda: self._left_to_rot(zero_stats),
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(side="left", padx=10)

    def _revive(self, zero_stats, mode):
        for stat in zero_stats:
            if stat == "Unconscious":
                continue
            if self.max_health[stat] > 0:
                self.stat_health[stat] = 1
                self.health_bars[stat].set(1)
                self.health_labels[stat].set(f"1/{self.max_health[stat]}")
                health_cell = HEALTH_CELL_MAP[stat]
                self._set_cell_value(self.stats_ws, health_cell, 1)
        self._save_all()
        self._update_abilities_state()
        if self.invalid_overlay:
            self.invalid_overlay.destroy()
            self.invalid_overlay = None

    def _left_to_rot(self, zero_stats):
        if self.invalid_overlay:
            self.invalid_overlay.destroy()
            self.invalid_overlay = None

        blackout = tk.Toplevel(self)
        blackout.configure(bg="black")
        blackout.attributes("-fullscreen", True)

        # Centered text
        label = tk.Label(
            blackout,
            text="You are left to rot...",
            bg="black",
            fg="red",
            font=("Papyrus", 48, "bold")
        )
        label.pack(expand=True)

        # Revive button (top-right corner)
        def revive_action():
            blackout.destroy()
            self._revive(zero_stats, "rot")

        tk.Button(
            blackout,
            text="Revive",
            command=revive_action,
            bg="white",
            fg="black",
            font=("Papyrus", 16, "bold")
        ).place(relx=0.98, rely=0.02, anchor="ne")

        # Quit button (bottom-left corner)
        tk.Button(
            blackout,
            text="Quit",
            command=self.destroy,
            bg="white",
            fg="black",
            font=("Papyrus", 16, "bold")
        ).place(relx=0.02, rely=0.98, anchor="sw")



    
    def _build_abilities_tab(self):
        for widget in self.abilities_frame.winfo_children():
            widget.destroy()

        header = tk.Frame(self.abilities_frame, bg=DARK_BG)
        header.pack(fill="x", pady=10)

        tk.Label(
            header,
            text="Abilities",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 14, "bold"),
        ).pack(side="left", padx=10)

        body = tk.Frame(self.abilities_frame, bg=DARK_BG)
        body.pack(fill="both", expand=True, padx=10, pady=10)

        self.abilities_buttons.clear()

        for stat in STATS:
            frame = tk.LabelFrame(body, text=stat, bg=DARK_BG, fg=DARK_FG)
            frame.pack(fill="x", pady=5)

            if stat == "Unconscious":
                tk.Label(
                    frame,
                    text="(No abilities)",
                    bg=DARK_BG,
                    fg="#888888",
                ).pack(padx=5, pady=5)
                continue

            abilities = self._get_abilities_for_stat(stat)
            if not abilities:
                tk.Label(
                    frame,
                    text="No abilities available.",
                    bg=DARK_BG,
                    fg="#888888",
                ).pack(padx=5, pady=5)
                continue

            for ability in abilities:
                name, func_text, cost_stars = ability
                cost = len(cost_stars.replace(" ", "")) if cost_stars else 0
                text = f"{name}  (Cost: {cost})"
                btn = tk.Button(
                    frame,
                    text=text,
                    bg="#1e1e1e",
                    fg=DARK_FG,
                    anchor="w",
                    command=lambda s=stat, n=name, c=cost: self._use_ability(s, n, c),
                )
                btn.pack(fill="x", padx=5, pady=2)
                Tooltip(btn, func_text or "No description.")
                self.abilities_buttons.append((stat, btn, cost))

        self._update_abilities_state()

    def _get_abilities_for_stat(self, stat):
        level = self.stat_levels.get(stat, 0)
        if level <= 1:
            return []

        layout = self.config["abilities_layout"]
        start_row = int(layout[f"{stat.lower()}_start_row"])

        # Each ability uses 3 rows: name, function, charges
        # Columns: D-F (level 2), G-I (level 3), J-L (level 4)
        col_groups = {
            2: ("D", "E", "F"),
            3: ("G", "H", "I"),
            4: ("J", "K", "L"),
        }

        abilities = []
        for lvl in range(2, level + 1):
            if lvl not in col_groups:
                continue
            cols = col_groups[lvl]
            name_cell = f"{cols[0]}{start_row}"
            func_cell = f"{cols[1]}{start_row+1}"
            cost_cell = f"{cols[2]}{start_row+2}"

            name = self._cell_value(self.abilities_ws, name_cell)
            func_text = self._cell_value(self.abilities_ws, func_cell)
            cost_stars = self._cell_value(self.abilities_ws, cost_cell)

            if name is None and func_text is None and cost_stars is None:
                continue

            abilities.append((str(name or "").strip(), str(func_text or "").strip(), str(cost_stars or "").strip()))

        return abilities

    def _update_abilities_state(self):
        for stat, btn, cost in self.abilities_buttons:
            enabled = True
            if stat in ["Body", "Mind", "Soul"]:
                if self.stat_health.get(stat, 0) <= 0:
                    enabled = False
            if cost > self.charges:
                enabled = False
            btn.configure(state="normal" if enabled else "disabled")

    def _use_ability(self, stat, name, cost):
        if cost > self.charges:
            messagebox.showwarning("Not enough charges", "You do not have enough charges.")
            return

        self.charges -= cost
        self.charges_var.set(self.charges)
        self._set_cell_value(self.abilities_ws, CHARGES_CELL, self.charges)
        self._save_all()
        self._update_abilities_state()

        messagebox.showinfo("Ability used", f"{name} has been used!")

    # ---------- New Item Tab ----------

    def _build_new_item_tab(self):
        frame = self.new_item_frame

        tk.Label(
            frame,
            text="NEW ITEM ACQUIRED!",
            bg=DARK_BG,
            fg=DARK_FG,
            font=("Papyrus", 14, "bold"),
        ).pack(pady=10)

        tk.Label(
            frame,
            text="Enter stat modifications for the new item.",
            bg=DARK_BG,
            fg=DARK_FG,
        ).pack(pady=5)

        tk.Button(
            frame,
            text="Open Item Creator",
            command=self._open_item_creator,
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(pady=5)

        form = tk.Frame(frame, bg=DARK_BG)
        form.pack(pady=10)

        # Health modifiers for Mind, Body, Soul
        tk.Label(form, text="Health Modifiers", bg=DARK_BG, fg=DARK_FG, font=("Papyrus", 10, "bold")).grid(
            row=0, column=0, columnspan=2, pady=(0, 5), sticky="w"
        )

        self.health_mod_vars = {}
        for i, stat in enumerate(["Body", "Mind", "Soul"], start=1):
            tk.Label(form, text=f"{stat} HP Δ:", bg=DARK_BG, fg=DARK_FG).grid(row=i, column=0, sticky="e", padx=5, pady=2)
            var = tk.StringVar()
            tk.Entry(form, textvariable=var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG, width=8).grid(
                row=i, column=1, sticky="w", padx=5, pady=2
            )
            self.health_mod_vars[stat] = var

        # Charges
        tk.Label(form, text="Charges Δ:", bg=DARK_BG, fg=DARK_FG).grid(
            row=4, column=0, sticky="e", padx=5, pady=(10, 2)
        )
        self.charges_mod_var = tk.StringVar()
        tk.Entry(form, textvariable=self.charges_mod_var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG, width=8).grid(
            row=4, column=1, sticky="w", padx=5, pady=(10, 2)
        )

        # Level modifiers for all stats
        tk.Label(form, text="Level Modifiers", bg=DARK_BG, fg=DARK_FG, font=("Papyrus", 10, "bold")).grid(
            row=5, column=0, columnspan=2, pady=(10, 5), sticky="w"
        )

        self.level_mod_vars = {}
        for i, stat in enumerate(STATS, start=6):
            tk.Label(form, text=f"{stat} Level Δ:", bg=DARK_BG, fg=DARK_FG).grid(
                row=i, column=0, sticky="e", padx=5, pady=2
            )
            var = tk.StringVar()
            tk.Entry(form, textvariable=var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG, width=8).grid(
                row=i, column=1, sticky="w", padx=5, pady=2
            )
            self.level_mod_vars[stat] = var

        tk.Button(
            frame,
            text="Apply Item Modifiers",
            command=self._apply_item_modifiers,
            bg="#1e1e1e",
            fg=DARK_FG,
        ).pack(pady=15)

    def _open_item_creator(self):
        creator = tk.Toplevel(self)
        creator.title("Create Item")
        creator.configure(bg=DARK_BG)

        tk.Label(creator, text="Item Name:", bg=DARK_BG, fg=DARK_FG).grid(row=0, column=0, sticky="e", padx=5, pady=2)
        name_var = tk.StringVar()
        tk.Entry(creator, textvariable=name_var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG).grid(row=0, column=1, padx=5, pady=2)

        # Basic stat modifiers
        stat_vars = {}
        for i, stat in enumerate(["Body", "Mind", "Soul", "Unconscious"], start=1):
            tk.Label(creator, text=f"{stat} Δ:", bg=DARK_BG, fg=DARK_FG).grid(row=i, column=0, sticky="e", padx=5, pady=2)
            var = tk.StringVar(value="0")
            tk.Entry(creator, textvariable=var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG, width=8).grid(row=i, column=1, sticky="w", padx=5, pady=2)
            stat_vars[stat.lower()] = var

        # Extra headers (key:value per line)
        tk.Label(creator, text="Extra headers (key:value per line):", bg=DARK_BG, fg=DARK_FG).grid(row=5, column=0, columnspan=2, sticky="w", padx=5, pady=(10, 2))
        extras_text = tk.Text(creator, height=6, width=40, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG)
        extras_text.grid(row=6, column=0, columnspan=2, padx=5, pady=2)

        # PNG selection
        tk.Label(creator, text="Optional PNG:", bg=DARK_BG, fg=DARK_FG).grid(row=7, column=0, sticky="e", padx=5, pady=(10, 2))
        png_path_var = tk.StringVar()
        tk.Entry(creator, textvariable=png_path_var, bg="#1e1e1e", fg=DARK_FG, insertbackground=DARK_FG, width=30).grid(row=7, column=1, sticky="w", padx=5, pady=(10, 2))

        def browse_png():
            from tkinter import filedialog
            path = filedialog.askopenfilename(filetypes=[("PNG files", "*.png")])
            if path:
                png_path_var.set(path)

        tk.Button(creator, text="Browse...", command=browse_png, bg="#1e1e1e", fg=DARK_FG).grid(row=7, column=2, padx=5, pady=(10, 2))

        def save_item():
            import shutil

            name = name_var.get().strip()
            if not name:
                messagebox.showerror("Error", "Item name is required.")
                return

            # Create folder
            items_root = "items"
            os.makedirs(items_root, exist_ok=True)
            item_folder = os.path.join(items_root, name)
            os.makedirs(item_folder, exist_ok=True)

            # Build data dict
            data = {"name": name}
            for key, var in stat_vars.items():
                try:
                    data[key] = int(var.get().strip() or "0")
                except ValueError:
                    messagebox.showerror("Error", f"Invalid value for {key}.")
                    return

            extra_lines = extras_text.get("1.0", "end-1c").splitlines()
            for line in extra_lines:
                if ":" in line:
                    k, v = line.split(":", 1)
                    data[k.strip()] = v.strip()

            # Write item.txt
            txt_path = os.path.join(item_folder, "item.txt")
            with open(txt_path, "w", encoding="utf-8") as f:
                for k, v in data.items():
                    f.write(f"{k}\t{v}\n")

            # Copy PNG if provided
            png_src = png_path_var.get().strip()
            if png_src and os.path.exists(png_src):
                png_dst = os.path.join(item_folder, os.path.basename(png_src))
                shutil.copy2(png_src, png_dst)

            messagebox.showinfo("Item Saved", f"Item '{name}' saved to {item_folder}.")
            creator.destroy()

        tk.Button(creator, text="Save Item", command=save_item, bg="#1e1e1e", fg=DARK_FG).grid(row=8, column=0, columnspan=3, pady=10)


    def _apply_item_modifiers(self):
        # Health mods
        for stat in ["Body", "Mind", "Soul"]:
            delta_str = self.health_mod_vars[stat].get().strip()
            if delta_str:
                try:
                    delta = int(delta_str)
                except ValueError:
                    messagebox.showerror("Error", f"Invalid health modifier for {stat}.")
                    return
                if stat in self.stat_health:
                    new_hp = max(0, self.stat_health[stat] + delta)
                    max_hp = self.max_health[stat]
                    new_hp = min(new_hp, max_hp)
                    self.stat_health[stat] = new_hp
                    self.health_bars[stat].set(new_hp)
                    self.health_labels[stat].set(f"{new_hp}/{max_hp}")
                    health_cell = HEALTH_CELL_MAP[stat]
                    self._set_cell_value(self.stats_ws, health_cell, new_hp)

        # Charges
        charges_delta_str = self.charges_mod_var.get().strip()
        if charges_delta_str:
            try:
                delta = int(charges_delta_str)
            except ValueError:
                messagebox.showerror("Error", "Invalid charges modifier.")
                return
            self.charges = max(0, self.charges + delta)
            self.charges_var.set(self.charges)
            self._set_cell_value(self.abilities_ws, CHARGES_CELL, self.charges)

        # Level mods
        for stat in STATS:
            delta_str = self.level_mod_vars[stat].get().strip()
            if delta_str:
                try:
                    delta = int(delta_str)
                except ValueError:
                    messagebox.showerror("Error", f"Invalid level modifier for {stat}.")
                    return
                new_level = max(0, min(4, self.stat_levels[stat] + delta))
                self.stat_levels[stat] = new_level
                level_cell = STAT_CELL_MAP[stat]
                self._set_cell_value(self.stats_ws, level_cell, new_level)

                # Recompute max HP for non-unconscious
                if stat != "Unconscious":
                    max_hp = 0
                    if new_level > 0:
                        max_hp = min(4, new_level) + 1
                    self.max_health[stat] = max_hp
                    # If current HP > new max, clamp
                    current = min(self.stat_health[stat], max_hp)
                    self.stat_health[stat] = current
                    self.health_bars[stat].set(current)
                    self.health_labels[stat].set(f"{current}/{max_hp}")
                    health_cell = HEALTH_CELL_MAP[stat]
                    self._set_cell_value(self.stats_ws, health_cell, current)

        self._save_all()
        self._build_abilities_tab()
        self._update_abilities_state()
        messagebox.showinfo("Item Applied", "New item modifiers have been applied.")

        # Clear fields
        for v in self.health_mod_vars.values():
            v.set("")
        self.charges_mod_var.set("")
        for v in self.level_mod_vars.values():
            v.set("")


if __name__ == "__main__":
    try:
        app = DnDApp()
        app.mainloop()
    except Exception as e:
        print("An error occurred before the GUI could start:", e)
